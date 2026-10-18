import openpyxl
import os


def process_xlsx_file(file_path):
    """
    Открывает xlsx-файл, ищет в ячейках текст '.cpp', если находит, удаляет всё от конца до первого встреченного '/'
    Также удаляет ячейки, содержащие текст 'Tests'
    """
    # Проверка доступности файла
    if not os.path.exists(file_path):
        print(f"Файл {file_path} не найден!")
        return

    try:
        # Открыть файл
        workbook = openpyxl.load_workbook(file_path)

        # Индикация работы на каждом листе
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            print(f"\nОбработка листа: {sheet_name}")

            cells_modified = 0
            cells_deleted = 0

            # Создание спика ячеек для обработки
            cells_to_process = list(sheet.iter_rows())

            for row in cells_to_process:
                for cell in row:
                    # Проверка на наличе текста 'Tests'
                    if isinstance(cell.value, str) and 'Tests' in cell.value:
                        cell.value = None  # Очистка ячейки
                        cells_deleted += 1
                        continue

                    # Проверка на наличе текста '.cpp'
                    if isinstance(cell.value, str) and '.cpp' in cell.value:
                        # Обнаружение последнего '/' перед '.cpp'
                        cpp_pos = cell.value.find('.cpp')
                        slash_pos = cell.value.rfind('/', 0, cpp_pos)

                        if slash_pos != -1:  # если '/' найден
                            # Удаление данных с конца
                            new_value = cell.value[:slash_pos]
                            cell.value = new_value
                            cells_modified += 1

            if cells_modified > 0:
                print(f"На листе '{sheet_name}' изменено {cells_modified} ячеек (обработка .cpp)")
            if cells_deleted > 0:
                print(f"На листе '{sheet_name}' очищено {cells_deleted} ячеек (содержали 'Tests')")
            if cells_modified == 0 and cells_deleted == 0:
                print(f"На листе '{sheet_name}' совпадений не найдено")

        # сохранение
        workbook.save(file_path)
        print(f"\nФайл {file_path} успешно обработан и сохранён!")

    except Exception as e:
        print(f"Ошибка при обработке файла: {e}")


file_path = "/Users/apetrova/Downloads/cpp.xlsx"
process_xlsx_file(file_path)
